"""Preview checked originals through authenticated HTTP, independent of RAG indexing."""

import hashlib
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import func, select

from enterprise_agent.config import PROJECT_ROOT
from enterprise_agent.models import KnowledgeChunk, KnowledgeDocument

pytestmark = [
    pytest.mark.integration,
    pytest.mark.postgres,
    pytest.mark.usefixtures("pg_seed"),
]
BASE = "/api/v1/knowledge/documents"


async def headers(client, username="demo_customer"):
    response = await client.post(
        "/api/v1/auth/login", json={"username": username, "password": "DemoOnly!2026"}
    )
    assert response.status_code == 200, response.text
    return {"Authorization": "Bearer " + response.json()["access_token"]}


async def original(pg_factory, settings, content, file_type="markdown", **changes):
    document_id = uuid4()
    relative = f"knowledge/uploads/{document_id}.{file_type}"
    path = settings.data_dir / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    document = KnowledgeDocument(
        id=document_id,
        title="文档原件预览",
        file_type=file_type,
        file_path=relative,
        checksum=hashlib.sha256(content).hexdigest(),
        version=3,
        visibility="public",
        status="uploaded",
        **changes,
    )
    async with pg_factory() as session, session.begin():
        session.add(document)
    return document, path


def format_bytes(file_type):
    if file_type == "pdf":
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {
                        NameObject("/F1"): DictionaryObject(
                            {
                                NameObject("/Type"): NameObject("/Font"),
                                NameObject("/Subtype"): NameObject("/Type1"),
                                NameObject("/BaseFont"): NameObject("/Helvetica"),
                            }
                        )
                    }
                )
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 40 700 Td (Complete original refund policy.) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        buffer = BytesIO()
        writer.write(buffer)
        return buffer.getvalue(), "Complete original refund policy."
    if file_type == "docx":
        document = Document()
        document.add_paragraph("完整售后政策")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "表格中的退款规则"
        document.add_paragraph("用户确认后提交")
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue(), "完整售后政策\n\n表格中的退款规则\n\n用户确认后提交"
    text = "# 完整售后政策\n第一段说明。\n\n退款申请需要审核。"
    return text.encode(), text


@pytest.mark.parametrize("file_type", ["pdf", "docx", "txt", "markdown"])
async def test_preview_four_real_formats_without_models_or_index(
    file_type, pg_api_client, pg_api_app, pg_factory
):
    settings = pg_api_app.state.settings
    assert settings.rag_enabled is False
    content, expected = format_bytes(file_type)
    document, _ = await original(pg_factory, settings, content, file_type)
    response = await pg_api_client.get(
        f"{BASE}/{document.id}/content", headers=await headers(pg_api_client)
    )
    assert response.status_code == 200, response.text
    assert response.json() == {
        "document_id": str(document.id),
        "title": document.title,
        "file_type": file_type,
        "version": 3,
        "content": expected,
    }
    async with pg_factory() as session:
        assert (
            await session.scalar(
                select(func.count())
                .select_from(KnowledgeChunk)
                .where(KnowledgeChunk.document_id == document.id)
            )
            == 0
        )


@pytest.mark.parametrize("status", ["uploaded", "indexing", "failed"])
async def test_preview_unindexed_original_is_complete_and_keeps_html_as_plain_text(
    status, pg_api_client, pg_api_app, pg_factory
):
    text = "# 售后政策\n" + "逐项说明完整保留。\n" * 120 + '<script>alert("preview")</script>\n结束'
    document, _ = await original(pg_factory, pg_api_app.state.settings, text.encode())
    async with pg_factory() as session, session.begin():
        stored = await session.get(KnowledgeDocument, document.id)
        stored.status = status
    response = await pg_api_client.get(
        f"{BASE}/{document.id}/content", headers=await headers(pg_api_client)
    )
    assert response.status_code == 200, response.text
    assert response.json()["content"] == text
    assert response.headers["content-type"] == "application/json"


async def test_preview_obeys_document_visibility_and_authentication(
    pg_api_client, pg_api_app, pg_factory
):
    document, _ = await original(
        pg_factory, pg_api_app.state.settings, "内部政策只允许客服和管理员查看".encode()
    )
    async with pg_factory() as session, session.begin():
        (await session.get(KnowledgeDocument, document.id)).visibility = "internal"
    url = f"{BASE}/{document.id}/content"
    assert (await pg_api_client.get(url)).status_code == 401
    for username in ("demo_customer", "demo_analyst", "demo_support", "demo_admin"):
        response = await pg_api_client.get(url, headers=await headers(pg_api_client, username))
        assert response.status_code == (200 if username in {"demo_support", "demo_admin"} else 404)
        if response.status_code == 404:
            assert "内部政策" not in response.text
            assert response.json()["error"]["code"] == "document_not_found"


async def test_preview_hides_deleted_and_missing_documents(pg_api_client, pg_api_app, pg_factory):
    document, _ = await original(
        pg_factory,
        pg_api_app.state.settings,
        "已删除原件".encode(),
        deleted_at=datetime.now(UTC),
    )
    auth = await headers(pg_api_client, "demo_admin")
    for identifier in (document.id, uuid4()):
        response = await pg_api_client.get(f"{BASE}/{identifier}/content", headers=auth)
        assert response.status_code == 404
        assert response.json()["error"]["code"] == "document_not_found"
        assert "已删除原件" not in response.text


@pytest.mark.parametrize(
    ("failure", "status", "code"),
    [
        ("missing", 409, "document_source_unavailable"),
        ("changed", 409, "document_source_changed"),
        ("malformed", 422, "document_parse_failed"),
        ("oversized", 413, "document_too_large"),
    ],
)
async def test_preview_rejects_unavailable_changed_invalid_or_oversized_originals_safely(
    failure, status, code, pg_api_client, pg_api_app, pg_factory
):
    settings = pg_api_app.state.settings
    content = b"private-document-fragment-not-a-pdf" if failure == "malformed" else b"policy"
    file_type = "pdf" if failure == "malformed" else "txt"
    if failure == "oversized":
        settings.rag_upload_max_bytes = 1024
        content = b"x" * 2048
    document, path = await original(pg_factory, settings, content, file_type)
    if failure == "missing":
        path.unlink()
    elif failure == "changed":
        path.write_bytes(b"changed-private-original")
    response = await pg_api_client.get(
        f"{BASE}/{document.id}/content", headers=await headers(pg_api_client)
    )
    assert response.status_code == status, response.text
    assert response.json()["error"]["code"] == code
    assert "private-document-fragment" not in response.text
    assert "changed-private-original" not in response.text
    assert str(settings.data_dir) not in response.text
    assert document.file_path not in response.text


async def registered_seed(pg_factory, content):
    document = KnowledgeDocument(
        id=uuid4(),
        title="已发布的种子文档原件",
        file_type="markdown",
        file_path="data/knowledge/after-sales.md",
        checksum=hashlib.sha256(content).hexdigest(),
        version=2,
        visibility="public",
        status="uploaded",
    )
    async with pg_factory() as session, session.begin():
        session.add(document)
    return document


async def test_seed_archive_preserves_exact_registered_original_when_checkout_has_changed(
    pg_api_client, pg_api_app, pg_factory
):
    content = "# 已发布的原件版本\n退款申请仍需审核。\n".encode()
    source = PROJECT_ROOT / "data/knowledge/after-sales.md"
    checkout_content = source.read_bytes()
    assert checkout_content != content
    document = await registered_seed(pg_factory, content)
    auth = await headers(pg_api_client)
    url = f"{BASE}/{document.id}"
    # Without an exact archived original, a changed checkout still fails closed.
    missing = await pg_api_client.get(url + "/content", headers=auth)
    assert missing.status_code == 409
    assert missing.json()["error"]["code"] == "document_source_changed"
    archive = pg_api_app.state.settings.data_dir / "knowledge-source-archive"
    archive.mkdir()
    (archive / (document.checksum + ".blob")).write_bytes(content)

    preview = await pg_api_client.get(url + "/content", headers=auth)
    download = await pg_api_client.get(url + "/download", headers=auth)
    assert preview.status_code == download.status_code == 200
    assert preview.json()["content"] == content.decode().strip()
    assert preview.json()["version"] == 2
    assert download.content == content
    assert source.read_bytes() == checkout_content
    async with pg_factory() as session:
        stored = await session.get(KnowledgeDocument, document.id)
        assert stored.checksum == hashlib.sha256(content).hexdigest()
        assert stored.version == 2


async def test_corrupt_seed_archive_never_falls_back_to_a_valid_checkout(
    pg_api_client, pg_api_app, pg_factory
):
    content = (PROJECT_ROOT / "data/knowledge/after-sales.md").read_bytes()
    document = await registered_seed(pg_factory, content)
    archive = pg_api_app.state.settings.data_dir / "knowledge-source-archive"
    archive.mkdir()
    (archive / (document.checksum + ".blob")).write_bytes(b"corrupted-private-archive")
    auth = await headers(pg_api_client)
    for operation in ("content", "download"):
        response = await pg_api_client.get(f"{BASE}/{document.id}/{operation}", headers=auth)
        assert response.status_code == 409
        assert response.json()["error"]["code"] == "document_source_changed"
        assert "corrupted-private-archive" not in response.text


async def test_seed_archive_cannot_escape_archive_directory_through_a_symlink(
    pg_api_client, pg_api_app, pg_factory, monkeypatch, request
):
    content = b"private-original-outside-the-archive"
    document = await registered_seed(pg_factory, content)
    settings = pg_api_app.state.settings
    outside = settings.data_dir / "outside-archive.blob"
    outside.write_bytes(content)
    archive_root = settings.data_dir / "knowledge-source-archive"
    archive_root.mkdir()
    archive = archive_root / (document.checksum + ".blob")
    try:
        archive.symlink_to(outside)
        request.node.user_properties.append(("archive_symlink_check", "real_filesystem_symlink"))
    except OSError:
        # Windows may deny unprivileged symlink creation. Exercise the same
        # resolved-path boundary with fault injection rather than skipping it.
        archive.write_bytes(content)
        original_resolve = Path.resolve

        def resolve(path, *args, **kwargs):
            return outside if path == archive else original_resolve(path, *args, **kwargs)

        monkeypatch.setattr(Path, "resolve", resolve)
        request.node.user_properties.append(("archive_symlink_check", "injected_resolved_path"))
    response = await pg_api_client.get(
        f"{BASE}/{document.id}/content", headers=await headers(pg_api_client)
    )
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "document_source_unavailable"
    assert "private-original-outside" not in response.text
    assert str(outside) not in response.text


async def test_seed_archive_is_never_used_for_uploaded_originals(
    pg_api_client, pg_api_app, pg_factory
):
    content = "上传原件始终使用自己的存储位置".encode()
    settings = pg_api_app.state.settings
    document, _ = await original(pg_factory, settings, content)
    archive = settings.data_dir / "knowledge-source-archive"
    archive.mkdir()
    (archive / (document.checksum + ".blob")).write_bytes(b"unrelated-archive-data")
    response = await pg_api_client.get(
        f"{BASE}/{document.id}/content", headers=await headers(pg_api_client)
    )
    assert response.status_code == 200, response.text
    assert response.json()["content"] == content.decode()
