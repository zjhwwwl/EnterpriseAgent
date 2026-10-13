"""Actual PDF, DOCX and UTF-8 bytes exercise extraction and source attribution."""

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pypdf import PdfWriter

from enterprise_agent.domain_errors import BusinessError
from enterprise_agent.rag import parsing
from enterprise_agent.rag.parsing import parse_document


def text_pdf(*pages: str) -> bytes:
    """Build a real small PDF with a Helvetica font, page streams and valid xref."""
    font_number = 3 + 2 * len(pages)
    kids = " ".join(f"{3 + 2 * index} 0 R" for index in range(len(pages)))
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode(),
    ]
    for index, page_text in enumerate(pages):
        escaped = page_text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
        content_number = 4 + 2 * index
        objects.extend(
            [
                (
                    f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    f"/Resources << /Font << /F1 {font_number} 0 R >> >> "
                    f"/Contents {content_number} 0 R >>"
                ).encode(),
                f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream",
            ]
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, value in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode() + value + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(output)


def docx_bytes(document) -> bytes:
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def rewrite_zip(content: bytes, *, replacements=None, extra=None) -> bytes:
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for member in source.infolist():
            target.writestr(
                member.filename,
                (replacements or {}).get(member.filename, source.read(member.filename)),
            )
        for name, value in (extra or {}).items():
            target.writestr(name, value)
    return output.getvalue()


def assert_rejected(content, file_type, *, code="document_parse_failed"):
    with pytest.raises(BusinessError) as caught:
        parse_document(content, file_type)
    assert caught.value.status_code == 422
    assert caught.value.code == code
    return caught.value


def test_real_pdf_extracts_each_page_and_preserves_one_based_page_numbers():
    sections = parse_document(
        text_pdf("Refund requests need confirmation.", "Delivery takes three business days."), "pdf"
    )
    assert [section.text for section in sections] == [
        "Refund requests need confirmation.",
        "Delivery takes three business days.",
    ]
    assert [section.source_location for section in sections] == [{"page": 1}, {"page": 2}]


def test_pdf_blank_pages_do_not_change_the_number_of_later_text_pages():
    sections = parse_document(text_pdf("", "A policy on page two."), "pdf")
    assert len(sections) == 1
    assert sections[0].text == "A policy on page two."
    assert sections[0].source_location == {"page": 2}


def test_blank_pdf_is_rejected_with_an_explicit_ocr_limitation():
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    output = BytesIO()
    writer.write(output)
    error = assert_rejected(output.getvalue(), "pdf")
    assert "OCR" in error.message
    assert "未提取到文本" in error.message


def test_encrypted_pdf_is_rejected_instead_of_guessing_a_password():
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.encrypt("fictional-test-password")
    output = BytesIO()
    writer.write(output)
    error = assert_rejected(output.getvalue(), "pdf")
    assert "加密" in error.message
    assert "fictional-test-password" not in error.message


def test_docx_preserves_body_and_table_order_with_traceable_locations():
    document = Document()
    document.add_paragraph("平台售后政策")
    document.add_paragraph("")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "七天无理由退货"
    table.cell(0, 1).text = "食品质量问题需提供说明"
    document.add_paragraph("退款申请需要用户确认")
    sections = parse_document(docx_bytes(document), "docx")
    assert [section.text for section in sections] == [
        "平台售后政策",
        "七天无理由退货",
        "食品质量问题需提供说明",
        "退款申请需要用户确认",
    ]
    assert sections[0].source_location == {"paragraph": 1}
    assert sections[1].source_location == {"table": 1, "row": 1, "cell": 1, "paragraph": 1}
    assert sections[2].source_location == {"table": 1, "row": 1, "cell": 2, "paragraph": 1}
    assert sections[3].source_location == {"paragraph": 3}


def test_merged_docx_cell_is_not_repeated_for_each_grid_position():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并单元格政策"
    sections = parse_document(docx_bytes(document), "docx")
    assert [section.text for section in sections] == ["合并单元格政策"]


def test_nested_docx_table_has_its_parent_path_in_the_location():
    document = Document()
    outer = document.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "嵌套表格配送政策"
    sections = parse_document(docx_bytes(document), "docx")
    assert len(sections) == 1
    assert sections[0].text == "嵌套表格配送政策"
    assert sections[0].source_location["table_path"] == "1/r1c1/t1"
    assert sections[0].source_location["table"] == 1


@pytest.mark.parametrize("file_type", ["txt", "markdown", "md"])
def test_utf8_bom_and_original_line_ranges_are_preserved(file_type):
    content = "\ufeff\r\n# 退款规则\r\n签收后七天内可申请。\r\n\r\n运费按政策处理。\r\n".encode()
    sections = parse_document(content, file_type)
    assert [section.text for section in sections] == [
        "# 退款规则\n签收后七天内可申请。",
        "运费按政策处理。",
    ]
    assert [section.source_location for section in sections] == [
        {"line_start": 2, "line_end": 3},
        {"line_start": 5, "line_end": 5},
    ]


@pytest.mark.parametrize(
    "content",
    [
        b"",
        b" \t\n ",
        "政策文本".encode("gbk"),
        b"valid\x00binary",
        b"valid\x1ccontrol",
        "乱码\ufffd".encode(),
        b"\xff\xfe\x00",
    ],
    ids=["empty", "blank", "gbk", "nul", "control-line-separator", "replacement", "binary"],
)
def test_empty_undecodable_or_binary_text_is_rejected(content):
    assert_rejected(content, "txt")


def test_wrong_formats_and_malformed_containers_do_not_echo_the_input():
    for content, file_type in [
        (b"secret-fragment-not-a-pdf", "pdf"),
        (b"%PDF-1.4\nsecret-fragment-broken", "pdf"),
        (b"secret-fragment-not-a-zip", "docx"),
        (text_pdf("A valid PDF."), "txt"),
        (docx_bytes(Document()), "markdown"),
    ]:
        error = assert_rejected(content, file_type)
        assert "secret-fragment" not in error.message


def test_empty_docx_and_truncated_zip_are_rejected():
    content = docx_bytes(Document())
    assert_rejected(content, "docx")
    assert_rejected(content[:-20], "docx")


def test_arbitrary_zip_is_not_accepted_as_docx():
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("notes.txt", "这是普通压缩包")
    assert_rejected(output.getvalue(), "docx")


def test_docx_rejects_path_traversal_and_macro_members():
    document = Document()
    document.add_paragraph("正常政策")
    content = docx_bytes(document)
    for name in ["../outside.xml", "/absolute.xml", "word/vbaProject.bin"]:
        corrupted = rewrite_zip(content, extra={name: b"unsafe"})
        assert_rejected(corrupted, "docx")


def test_docx_rejects_duplicate_archive_names():
    document = Document()
    document.add_paragraph("正常政策")
    content = docx_bytes(document)
    with pytest.warns(UserWarning, match="Duplicate name"):
        duplicated = rewrite_zip(content, extra={"word/document.xml": b"duplicate"})
    assert_rejected(duplicated, "docx")


def test_docx_rejects_dtd_entities_before_wordprocessing_xml_is_parsed():
    document = Document()
    document.add_paragraph("正常政策")
    content = docx_bytes(document)
    with ZipFile(BytesIO(content)) as archive:
        body = archive.read("word/document.xml")
    body = body.replace(
        b"<w:document",
        b'<!DOCTYPE document [<!ENTITY policy "injected">]><w:document',
        1,
    )
    error = assert_rejected(rewrite_zip(content, replacements={"word/document.xml": body}), "docx")
    assert "DTD" in error.message


def test_docx_rejects_an_incorrect_main_part_content_type():
    document = Document()
    document.add_paragraph("正常政策")
    content = docx_bytes(document)
    with ZipFile(BytesIO(content)) as archive:
        types = archive.read("[Content_Types].xml")
    types = types.replace(
        b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
        b"application/vnd.ms-word.document.macroEnabled.main+xml",
    )
    assert_rejected(rewrite_zip(content, replacements={"[Content_Types].xml": types}), "docx")


@pytest.mark.parametrize(
    ("limit_name", "limit"),
    [("MAX_DOCX_MEMBERS", 2), ("MAX_DOCX_EXPANDED_BYTES", 128), ("MAX_DOCX_XML_BYTES", 128)],
)
def test_real_docx_archive_cannot_exceed_configured_resource_bounds(monkeypatch, limit_name, limit):
    document = Document()
    document.add_paragraph("正常政策")
    content = docx_bytes(document)
    monkeypatch.setattr(parsing, limit_name, limit)
    assert_rejected(content, "docx", code="document_too_large")


def test_input_extracted_text_and_pdf_page_limits_are_enforced(monkeypatch):
    with monkeypatch.context() as changes:
        changes.setattr(parsing, "MAX_INPUT_BYTES", 3)
        assert_rejected(b"policy", "txt", code="document_too_large")
    with monkeypatch.context() as changes:
        changes.setattr(parsing, "MAX_EXTRACTED_CHARS", 3)
        assert_rejected("售后政策文本".encode(), "txt", code="document_too_large")
    with monkeypatch.context() as changes:
        changes.setattr(parsing, "MAX_PDF_PAGES", 1)
        assert_rejected(text_pdf("First page.", "Second page."), "pdf", code="document_too_large")


def test_unsupported_document_type_is_explicit():
    assert_rejected(b"presentation", "pptx", code="unsupported_document_type")
